import os
import sys
import json
import mimetypes
from pathlib import Path
from dotenv import load_dotenv
from google import genai
from google.genai import types
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

# 상위 디렉토리의 .env 로드
env_path = Path(__file__).parent.parent / '.env'
load_dotenv(env_path)

GEMINI_API_KEY = os.getenv('GEMINI_API_KEY')
MODEL = 'gemini-2.5-flash'

ANALYSIS_PROMPT = """
다음 메뉴판 이미지를 분석하여 메뉴 목록을 JSON 형식으로 추출해주세요.

응답 형식:
{
  "restaurant_name": "식당 이름 (없으면 알 수 없음)",
  "menus": [
    {
      "name": "메뉴명",
      "price": "가격 (없으면 가격 미표기)",
      "description": "원본 설명 (없으면 null)",
      "fun_description": "재밌고 센스있는 설명 (2-3문장, 과하지 않게)"
    }
  ]
}

fun_description 작성 지침:
- 메뉴의 특징을 살리되 과장 없이 자연스럽게
- 먹고 싶은 마음이 생기도록 감각적으로 표현
- 2-3문장 이내

반드시 유효한 JSON만 응답하세요. 코드블록(```) 없이 순수 JSON만.
"""


def get_mime_type(image_path: str) -> str:
    mime_type, _ = mimetypes.guess_type(image_path)
    if mime_type and mime_type.startswith('image/'):
        return mime_type
    ext = Path(image_path).suffix.lower().lstrip('.')
    fallback = {'jpg': 'image/jpeg', 'jpeg': 'image/jpeg', 'png': 'image/png',
                'gif': 'image/gif', 'webp': 'image/webp'}
    return fallback.get(ext, 'image/jpeg')


def analyze_menu(image_path: str) -> dict:
    client = genai.Client(api_key=GEMINI_API_KEY)

    with open(image_path, 'rb') as f:
        image_data = f.read()

    mime_type = get_mime_type(image_path)

    response = client.models.generate_content(
        model=MODEL,
        config=types.GenerateContentConfig(response_mime_type='application/json'),
        contents=[
            types.Part.from_bytes(data=image_data, mime_type=mime_type),
            ANALYSIS_PROMPT,
        ]
    )

    try:
        return json.loads(response.text)
    except json.JSONDecodeError as e:
        raise ValueError(f"Gemini 응답을 JSON으로 파싱 실패: {e}\n응답 내용: {response.text[:200]}")


def create_docx(menu_data: dict, output_path: str):
    doc = Document()

    # 제목
    title = doc.add_heading('메뉴 설명서', 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # 식당명
    restaurant = menu_data.get('restaurant_name', '알 수 없음')
    if restaurant != '알 수 없음':
        subtitle = doc.add_paragraph(f"\U0001F37D️  {restaurant}")
        subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = subtitle.runs[0]
        run.font.size = Pt(14)
        run.bold = True

    doc.add_paragraph()

    menus = menu_data.get('menus', [])
    for i, menu in enumerate(menus, 1):
        name = menu.get('name', '(이름 없음)')
        price = menu.get('price', '가격 미표기')
        fun_desc = menu.get('fun_description', '')

        # 메뉴명 헤딩
        doc.add_heading(f"{i}. {name}", 2)

        # 가격
        price_para = doc.add_paragraph()
        price_run = price_para.add_run(f"\U0001F4B0 가격: {price}")
        price_run.font.size = Pt(11)
        price_run.bold = True

        # 재밌는 설명
        fun_para = doc.add_paragraph()
        fun_run = fun_para.add_run(fun_desc)
        fun_run.font.size = Pt(10)

        # 원본 설명 (있는 경우)
        if menu.get('description'):
            orig_para = doc.add_paragraph()
            orig_run = orig_para.add_run(f"[원문] {menu['description']}")
            orig_run.font.size = Pt(9)
            orig_run.font.color.rgb = RGBColor(0x80, 0x80, 0x80)

        doc.add_paragraph()

    doc.save(output_path)
    return len(menus)


def main():
    if not GEMINI_API_KEY:
        print("오류: GEMINI_API_KEY가 설정되지 않았습니다.")
        print(f"  → {env_path} 파일에 GEMINI_API_KEY=<your_key> 추가 후 재실행")
        sys.exit(1)

    # 이미지 경로 입력
    if len(sys.argv) > 1:
        image_path = sys.argv[1]
    else:
        image_path = input("이미지 경로를 입력하세요: ").strip()

    image_path = image_path.strip('"').strip("'")

    if not os.path.exists(image_path):
        print(f"오류: 파일을 찾을 수 없습니다 → {image_path}")
        sys.exit(1)

    print(f"\n이미지 분석 중: {image_path}")
    print("  Gemini Vision API 호출 중...")

    menu_data = analyze_menu(image_path)

    # 출력 경로 결정
    image_stem = Path(image_path).stem
    results_dir = Path(__file__).parent / 'results'
    results_dir.mkdir(exist_ok=True)
    output_path = results_dir / f"{image_stem}.docx"

    count = create_docx(menu_data, str(output_path))

    print(f"\n완료! 설명서 저장: {output_path}")
    print(f"발견된 메뉴 {count}개:")
    for menu in menu_data.get('menus', []):
        print(f"  - {menu.get('name', '?')}  ({menu.get('price', '?')})")


if __name__ == '__main__':
    main()
